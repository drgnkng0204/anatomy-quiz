import openpyxl, json, re
from pathlib import Path
d=Path(__file__).parent
old=json.loads((d/'data.json').read_text(encoding='utf-8'))
secs=old['secs']
wb=openpyxl.load_workbook(d.parent/'解剖学_過去問対応表.xlsx',data_only=True);ws=wb.active
rows=[]
for r in range(2,ws.max_row+1):
    v=[ws.cell(row=r,column=c).value for c in range(1,10)]
    if v[0] is None: continue
    sec=str(v[7]).strip() if v[7] is not None else ''
    rows.append(dict(id=str(v[0]).strip(),q=v[1] or '',ja=v[2] or '',la=v[3] or '',yr=str(v[4] or ''),
        sib=[s.strip() for s in str(v[5] or '').split(',') if s.strip()],page=str(v[6] or ''),sec=sec))
ids={x['id'] for x in rows}
print('rows',len(rows),'missing sib',sorted({s for x in rows for s in x['sib'] if s not in ids})[:5])
secids={s['id'] for s in secs}
print('bad sec',[x['id']+':'+x['sec'] for x in rows if x['sec'].split('-')[0] not in secids])
oldq={x['id']:x for x in old['q']}
ch={k:[] for k in ('q','ja','la','sib','page','sec')}
for x in rows:
    o=oldq.get(x['id'])
    if not o: print('new',x['id']);continue
    for k in ch:
        if o[k]!=x[k]: ch[k].append(x['id'])
print({k:len(v) for k,v in ch.items()}); print('removed',[k for k in oldq if k not in ids])
for k in ('ja','la','q'): print(k,ch[k][:15])
import unicodedata
def nz(t):
    t=unicodedata.normalize('NFKC',t or '').lower()
    t=re.sub(r'\*.*','',t)
    t=t.replace('旁','傍').replace('靱','靭').replace('第三','第3')
    return re.sub(r'[\s・.,、。]','',t)
def toks(a):
    a=unicodedata.normalize('NFKC',a or '').replace('旁','傍').replace('靱','靭').replace('第三','第3')
    a=re.sub(r'\*.*','',a)
    parts=re.split(r'[()（）]',a)
    return {re.sub(r'[\s・.,、。]','',p).lower() for p in parts if re.sub(r'[\s・.,、。]','',p)}
def blanks(ja):
    m={}
    for l in (ja or '').split('\n'):
        x=re.match(r'^\s*([A-Z])\s*[:：]\s*(.*)$',l)
        if x:m[x.group(1)]=toks(x.group(2))
    return m
def equiv(a,b):
    x,y=blanks(a),blanks(b)
    if not x or set(x)!=set(y):return nz(a)==nz(b)
    return all(x[k]&y[k] for k in x)
def key(i):
    m=re.match(r'^(\d+)(再?)-(\d+)([a-z]?)$',i);return (int(m.group(1)),1 if m.group(2) else 0,int(m.group(3)),m.group(4))
byq={}
for x in rows: byq.setdefault(re.sub(r'\s+','',x['q']).lower(),[]).append(x)
par={x['id']:x['id'] for x in rows}
def find(a):
    while par[a]!=a:par[a]=par[par[a]];a=par[a]
    return a
for v in byq.values():
    for i in range(len(v)):
        for j in range(i+1,len(v)):
            if equiv(v[i]['ja'],v[j]['ja']):par[find(v[j]['id'])]=find(v[i]['id'])
grp={}
for x in rows: grp.setdefault(find(x['id']),[]).append(x['id'])
for ids in grp.values():
    c=min(ids,key=key)
    for i in ids: next(x for x in rows if x['id']==i)['c']=c
print('unique questions',len(grp),'merged-away',len(rows)-len(grp))
json.dump({'q':rows,'secs':secs},open(d/'data.json','w',encoding='utf-8'),ensure_ascii=False,separators=(',',':'))
